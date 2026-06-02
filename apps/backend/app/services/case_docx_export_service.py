"""Export case brief and case documents to Word (.docx) for printing and filing."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from uuid import UUID

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.documents import Document as CaseDocument
from app.models.legal import Case
from app.services.document_file_storage import DocumentFileStorage
from app.services.ocr_service import OcrService

BRIEF_SECTIONS: list[tuple[str, str]] = [
    ("statement_of_facts", "Statement of facts"),
    ("legal_issues", "Legal issues"),
    ("client_objectives", "Client objectives"),
    ("requested_legal_action", "Requested legal action"),
    ("special_instructions", "Special instructions"),
]

# Categories treated as pleadings / court papers in the export pack
PLEADING_CATEGORY_NAMES = frozenset(
    {
        "COURT_FILING",
        "AFFIDAVIT",
        "CONTRACT",
        "CORRESPONDENCE",
        "EVIDENCE",
    }
)


class CaseDocxExportService:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.storage = DocumentFileStorage()
        self.ocr = OcrService()

    async def get_case(self, case_id: UUID) -> Case | None:
        result = await self.db.execute(select(Case).where(Case.id == case_id, Case.deleted_at.is_(None)))
        return result.scalar_one_or_none()

    async def list_case_documents(self, case_id: UUID) -> list[CaseDocument]:
        result = await self.db.execute(
            select(CaseDocument)
            .where(CaseDocument.case_id == case_id, CaseDocument.deleted_at.is_(None))
            .options(selectinload(CaseDocument.category))
            .order_by(CaseDocument.created_at.asc())
        )
        return list(result.scalars().all())

    def _case_details(self, case: Case) -> dict:
        master = case.master_data or {}
        if isinstance(master, dict):
            details = master.get("case_details")
            if isinstance(details, dict):
                return details
        return {}

    def _letterhead(self, doc: Document, title: str, case: Case) -> None:
        firm = doc.add_paragraph()
        firm.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = firm.add_run("SANSON Law Firm")
        run.bold = True
        run.font.size = Pt(16)

        sub = doc.add_paragraph()
        sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        sub.add_run("SANSON Legal OS").font.size = Pt(10)

        doc.add_paragraph()
        doc.add_heading(title, level=1)
        doc.add_paragraph(f"Case number: {case.case_number}")
        doc.add_paragraph(f"Matter: {case.title}")
        doc.add_paragraph(f"Category: {case.case_category.value if hasattr(case.case_category, 'value') else case.case_category}")
        doc.add_paragraph(
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
        )
        doc.add_paragraph()

    def _add_brief_body(self, doc: Document, case: Case) -> None:
        details = self._case_details(case)
        if case.description and str(case.description).strip():
            doc.add_heading("Case summary", level=2)
            doc.add_paragraph(str(case.description).strip())

        for key, label in BRIEF_SECTIONS:
            raw = details.get(key)
            if key == "requested_legal_action" and not raw:
                raw = details.get("requested_action")
            text = str(raw).strip() if raw else ""
            if not text:
                continue
            doc.add_heading(label, level=2)
            doc.add_paragraph(text)

        if not case.description and not any(
            details.get(k) or details.get("requested_action")
            for k, _ in BRIEF_SECTIONS
        ):
            doc.add_paragraph("No case brief content recorded yet.")

    def build_brief_docx(self, case: Case) -> bytes:
        doc = Document()
        self._letterhead(doc, "Case Brief", case)
        self._add_brief_body(doc, case)
        return self._save(doc)

    async def _append_document_content(
        self, doc: Document, item: CaseDocument, include_images: bool = True
    ) -> None:
        cat_name = item.category.name if item.category else "OTHER"
        cat_label = item.category.display_name if item.category else "Document"
        doc.add_heading(f"{item.original_file_name or item.file_name}", level=2)
        doc.add_paragraph(f"Category: {cat_label} ({cat_name})")
        doc.add_paragraph(
            f"Uploaded: {item.created_at.strftime('%Y-%m-%d') if item.created_at else '—'}"
        )

        try:
            data = await self.storage.download_bytes(item.storage_path)
        except ValueError as exc:
            doc.add_paragraph(f"[Could not load file: {exc}]")
            doc.add_paragraph()
            return

        filename = item.original_file_name or item.file_name or "document"
        mime = item.mime_type or "application/octet-stream"
        lower = filename.lower()

        if lower.endswith(".docx") or "wordprocessingml" in mime:
            try:
                from docx import Document as DocxDocument

                source = DocxDocument(io.BytesIO(data))
                for para in source.paragraphs:
                    if para.text.strip():
                        doc.add_paragraph(para.text.strip())
            except Exception as exc:
                doc.add_paragraph(f"[Could not read Word file: {exc}]")
            doc.add_paragraph()
            return

        if lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")) and include_images:
            try:
                doc.add_picture(io.BytesIO(data), width=Inches(5.5))
                doc.add_paragraph()
            except Exception:
                text, _, _ = self.ocr.extract_text(data, mime, filename)
                if text.strip():
                    doc.add_paragraph(text.strip())
                else:
                    doc.add_paragraph("[Image — open original file in the case folder]")
            doc.add_paragraph()
            return

        text, confidence, meta = self.ocr.extract_text(data, mime, filename)
        if text.strip():
            doc.add_paragraph(text.strip())
            if confidence < 70:
                doc.add_paragraph(
                    f"(Extracted text — confidence {confidence:.0f}%; verify against original.)"
                )
        else:
            err = meta.get("error", "unsupported_type")
            doc.add_paragraph(
                f"[No text extracted from this file ({err}). Download the original from the case file list.]"
            )
        doc.add_paragraph()

    async def build_pleadings_docx(
        self, case: Case, documents: list[CaseDocument], *, all_documents: bool = False
    ) -> bytes:
        doc = Document()
        self._letterhead(doc, "Case Pleadings & Documents", case)

        doc.add_heading("Part I — Case brief", level=1)
        self._add_brief_body(doc, case)

        doc.add_page_break()
        doc.add_heading("Part II — Pleadings and case documents", level=1)

        if all_documents:
            items = documents
        else:
            items = [
                d
                for d in documents
                if d.category and d.category.name in PLEADING_CATEGORY_NAMES
            ]
            if not items:
                items = documents

        if not items:
            doc.add_paragraph("No documents uploaded for this case yet.")
        else:
            for item in items:
                await self._append_document_content(doc, item)

        return self._save(doc)

    @staticmethod
    def _save(doc: Document) -> bytes:
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def safe_filename(self, case_number: str, suffix: str) -> str:
        safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in case_number)[:80]
        return f"{safe}_{suffix}.docx"
